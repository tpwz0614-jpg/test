import datetime as dt

import openpyxl
import pytest

from irr_xirr_tool import io as cfio


class TestManual:
    def test_parse_manual_amounts(self):
        s = cfio.parse_manual_amounts(["-1000000", "300000", "300,000", 300000])
        assert s.amounts == [-1000000.0, 300000.0, 300000.0, 300000.0]
        assert not s.is_dated

    def test_parse_manual_dated(self):
        s = cfio.parse_manual_dated([("2020-01-01", -1000), ("2020-07-01", 1200)])
        assert s.is_dated
        assert s.dates == [dt.date(2020, 1, 1), dt.date(2020, 7, 1)]
        assert s.amounts == [-1000.0, 1200.0]

    def test_bad_amount_raises(self):
        with pytest.raises(ValueError):
            cfio.parse_manual_amounts(["-1000", "not a number"])


class TestCsv:
    def test_irr_csv_no_header(self, tmp_path):
        p = tmp_path / "cf.csv"
        p.write_text("-1000000\n300000\n300000\n300000\n300000\n300000\n")
        s = cfio.read_csv(p)
        assert not s.is_dated
        assert s.amounts == [-1000000.0, 300000.0, 300000.0, 300000.0, 300000.0, 300000.0]

    def test_irr_csv_with_header(self, tmp_path):
        p = tmp_path / "cf.csv"
        p.write_text("amount\n-1000000\n300000\n")
        s = cfio.read_csv(p)
        assert not s.is_dated
        assert s.amounts == [-1000000.0, 300000.0]

    def test_xirr_csv_with_header(self, tmp_path):
        p = tmp_path / "cf.csv"
        p.write_text("date,amount\n2020-01-01,-10000\n2020-07-01,2750\n2021-01-01,9000\n")
        s = cfio.read_csv(p)
        assert s.is_dated
        assert s.dates == [dt.date(2020, 1, 1), dt.date(2020, 7, 1), dt.date(2021, 1, 1)]
        assert s.amounts == [-10000.0, 2750.0, 9000.0]

    def test_xirr_csv_no_header_date_first(self, tmp_path):
        p = tmp_path / "cf.csv"
        p.write_text("2020-01-01,-10000\n2020-07-01,2750\n")
        s = cfio.read_csv(p)
        assert s.is_dated

    def test_comma_thousands_separator(self, tmp_path):
        p = tmp_path / "cf.csv"
        p.write_text("amount\n\"-1,000,000\"\n\"300,000\"\n")
        s = cfio.read_csv(p)
        assert s.amounts == [-1000000.0, 300000.0]

    def test_malformed_row_raises_with_row_number(self, tmp_path):
        p = tmp_path / "cf.csv"
        p.write_text("amount\n-1000\nnot-a-number\n")
        with pytest.raises(ValueError, match="Row 2"):
            cfio.read_csv(p)


class TestExcel:
    def test_irr_excel(self, tmp_path):
        p = tmp_path / "cf.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["amount"])
        for v in (-1_000_000, 300000, 300000, 300000, 300000, 300000):
            ws.append([v])
        wb.save(p)
        s = cfio.read_excel(p)
        assert not s.is_dated
        assert s.amounts[0] == -1_000_000.0
        assert len(s.amounts) == 6

    def test_xirr_excel_native_dates(self, tmp_path):
        p = tmp_path / "cf.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["date", "amount"])
        ws.append([dt.date(2020, 1, 1), -10000])
        ws.append([dt.date(2020, 7, 1), 2750])
        ws.append([dt.date(2021, 1, 1), 9000])
        wb.save(p)
        s = cfio.read_excel(p)
        assert s.is_dated
        assert s.dates == [dt.date(2020, 1, 1), dt.date(2020, 7, 1), dt.date(2021, 1, 1)]

    def test_load_cashflows_dispatch(self, tmp_path):
        p = tmp_path / "cf.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["amount"])
        ws.append([-1000])
        ws.append([1200])
        wb.save(p)
        s = cfio.load_cashflows(p)
        assert s.amounts == [-1000.0, 1200.0]

    def test_load_cashflows_unsupported_extension(self, tmp_path):
        p = tmp_path / "cf.txt"
        p.write_text("x")
        with pytest.raises(ValueError, match="Unsupported"):
            cfio.load_cashflows(p)
